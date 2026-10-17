"""캡션 시험. 시험용 pptx를 코드로 만들어 그림 꺼내기, 건너뛰기, 저장, 작업 스레드를 본다."""
from __future__ import annotations

import io
import threading

import pytest
from PIL import Image
from pptx import Presentation
from pptx.util import Inches

from ppt_reader import captions as cap
from ppt_reader.captions import CaptionStore, CaptionWorker, extract, parse_captions, plan, to_llm_image
from ppt_reader.llm.base import LLMError


def png(w, h, color=(200, 80, 60), alpha=False) -> bytes:
    im = Image.new("RGBA" if alpha else "RGB", (w, h), color + ((128,) if alpha else ()))
    im.putpixel((0, 0), (1, 2, 3, 255) if alpha else (1, 2, 3))  # 그림마다 바이트가 달라지게
    out = io.BytesIO()
    im.save(out, "PNG")
    return out.getvalue()


def add_pic(slide, data, left=1, top=1, width=3, height=2):
    return slide.shapes.add_picture(io.BytesIO(data), Inches(left), Inches(top), Inches(width), Inches(height))


@pytest.fixture
def deck(tmp_path):
    prs = Presentation()
    blank = prs.slide_layouts[6]
    s1 = prs.slides.add_slide(blank)  # 1: 그림 하나 + 그룹 안 그림 둘 + 잘라낸 그림
    add_pic(s1, png(200, 100, (10, 20, 30)))
    group = s1.shapes.add_group_shape()
    add_pic(group, png(120, 120, (40, 50, 60)), 4, 1)
    add_pic(group, png(130, 90, (70, 80, 90)), 4, 3)
    cropped = add_pic(s1, png(400, 200, (100, 110, 120)), 1, 4)
    cropped.crop_left, cropped.crop_right = 0.25, 0.25
    s2 = prs.slides.add_slide(prs.slide_layouts[8])  # 2: 그림 자리표시자
    holder = next(ph for ph in s2.placeholders if ph.placeholder_format.type == 18)  # PICTURE
    holder.insert_picture(io.BytesIO(png(300, 200, (130, 140, 150))))
    s3 = prs.slides.add_slide(blank)  # 3: 아주 작은 그림 두 가지
    add_pic(s3, png(32, 32, (160, 170, 180)))  # 원본이 작음
    add_pic(s3, png(500, 500, (190, 200, 210)), width=0.3, height=0.3)  # 슬라이드 위에서 작음
    logo = png(100, 100, (5, 5, 5))
    for _ in range(4):  # 4~7: 같은 로고가 네 장에
        add_pic(prs.slides.add_slide(blank), logo, 8, 0.2, 1, 1)
    path = tmp_path / "그림 시험.pptx"
    prs.save(path)
    return path


def test_extract_groups_placeholders_crop(deck):
    slides = extract(Presentation(deck))
    assert [len(slides[n]) for n in range(1, 8)] == [4, 1, 2, 1, 1, 1, 1]
    s1 = slides[1]
    assert [p.index for p in s1] == [1, 2, 3, 4] and all(not p.skip for p in s1)  # 그룹 안 그림도
    assert s1[3].crop == (0.25, 0.0, 0.25, 0.0) and s1[3].key.endswith(":0.2500,0.0000,0.2500,0.0000")
    assert slides[2][0].content_type == "image/png" and not slides[2][0].skip  # 자리표시자 그림
    assert [p.skip for p in slides[3]] == [cap.SKIP_TINY, cap.SKIP_TINY]
    assert all(slides[n][0].skip == cap.SKIP_REPEAT for n in range(4, 8))  # 세 장 넘게 되풀이되면 로고
    todo, over = plan(slides)
    assert len(todo) == 5 and not over


def test_three_repeats_are_kept(tmp_path):
    prs = Presentation()
    logo = png(100, 100)
    for _ in range(3):
        add_pic(prs.slides.add_slide(prs.slide_layouts[6]), logo)
    path = tmp_path / "세 번.pptx"
    prs.save(path)
    assert all(not pics[0].skip for pics in extract(Presentation(path)).values())


def test_limit_40(tmp_path):
    prs = Presentation()
    for i in range(41):
        add_pic(prs.slides.add_slide(prs.slide_layouts[6]), png(80, 80, (i, i, i)))
    path = tmp_path / "많음.pptx"
    prs.save(path)
    todo, over = plan(extract(Presentation(path)))
    assert len(todo) == 40 and over and todo[-1].slide == 40


def test_to_llm_image_resizes_crops_and_formats(deck):
    big = cap.SlidePicture(1, 1, blob=png(2000, 1000), content_type="image/png")
    im = Image.open(io.BytesIO(to_llm_image(big)))
    assert im.format == "JPEG" and im.size == (1024, 512)  # 투명이 없으면 JPEG, 긴 변 1024
    cropped = cap.SlidePicture(1, 1, blob=png(400, 200, alpha=True), content_type="image/png", crop=(0.25, 0, 0.25, -0.1))
    im = Image.open(io.BytesIO(to_llm_image(cropped)))
    assert im.format == "PNG" and im.size == (200, 200)  # 투명이 있으면 PNG, 음수 잘라내기는 0


def test_parse_captions():
    text = "그림 1: 막대그래프로 성적을 보여 준다.\n잡담\n그림 2 : 학교 사진.\n그림 9: 없는 그림"
    assert parse_captions(text, 2) == {1: "막대그래프로 성적을 보여 준다.", 2: "학교 사진."}


def test_store_keys_survive_rename(deck, tmp_path, app_home):
    store = CaptionStore(app_home)
    slides = extract(Presentation(deck))
    store.put(slides[1][0].key, "첫 그림 설명", "claude-haiku-5-5")
    moved = tmp_path / "옮긴 파일.pptx"
    moved.write_bytes(deck.read_bytes())
    again = CaptionStore(app_home)
    texts = again.texts(extract(Presentation(moved)))
    assert texts[1][0] == "첫 그림 설명" and texts[1][1] == ""  # 파일 이름을 바꿔도 다시 쓴다
    again.edit(slides[1][0].key, "고친 설명")
    assert CaptionStore(app_home).get(slides[1][0].key)["edited"]


class FakeProvider:
    def __init__(self, fail=None):
        self.requests = []
        self.fail = fail

    def stream(self, req, cancel):
        self.requests.append(req)
        if self.fail:
            raise self.fail
        yield "\n".join(f"그림 {i}: 설명 {len(self.requests)}-{i}" for i in range(1, len(req.images) + 1))


def run_worker(store, provider, todo, texts=None, redo=False):
    progress, failures, done = [], [], threading.Event()

    def on_progress(d, total):
        progress.append((d, total))
        if d == total:
            done.set()

    def on_failed(e):
        failures.append(e.kind)
        done.set()

    w = CaptionWorker(store, lambda: provider, on_progress, on_failed)
    w.start("claude-haiku-5-5", texts or {}, todo, redo)
    assert done.wait(5)
    w.close()
    return progress, failures


def test_worker_makes_and_keeps_edited(deck, app_home):
    store = CaptionStore(app_home)
    slides = extract(Presentation(deck))
    todo, _ = plan(slides)
    provider = FakeProvider()
    progress, failures = run_worker(store, provider, todo, {1: "슬라이드 1 글"})
    assert failures == [] and progress[0] == (0, 5) and progress[-1] == (5, 5)
    assert len(provider.requests) == 2  # 장마다 한 번(1장 그림 4개, 2장 그림 1개)
    req = provider.requests[0]
    assert req.effort == "low" and len(req.images) == 4 and "슬라이드 1 글" in req.question
    assert store.get(todo[0].key)["text"] == "설명 1-1"
    store.edit(todo[0].key, "발표자가 고친 설명")
    provider2 = FakeProvider()
    run_worker(store, provider2, todo, redo=False)
    assert provider2.requests == []  # 이미 있는 캡션은 다시 만들지 않는다
    run_worker(store, provider2, todo, redo=True)  # [다시 만들기]
    assert store.get(todo[0].key)["text"] == "발표자가 고친 설명"  # 고친 캡션은 덮어쓰지 않는다
    assert store.get(todo[1].key)["text"].startswith("설명 1-")


def test_worker_reports_failure(deck, app_home):
    todo, _ = plan(extract(Presentation(deck)))
    _, failures = run_worker(CaptionStore(app_home), FakeProvider(fail=LLMError("key", "bad")), todo)
    assert failures == ["key"]
