import sys
from pathlib import Path

import pytest
from pptx import Presentation
from pptx.util import Inches

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

from ppt_reader.link import resolve_local_path  # noqa: E402
from ppt_reader.slides import Slide, load_deck, to_speech  # noqa: E402


@pytest.fixture
def deck_path(tmp_path):
    prs = Presentation()
    s = prs.slides.add_slide(prs.slide_layouts[1])
    s.shapes.title.text = "제목"
    s.placeholders[1].text = "첫 줄\n둘째 줄"
    s.notes_slide.notes_text_frame.text = "노트 대본입니다."

    s = prs.slides.add_slide(prs.slide_layouts[6])
    # 일부러 오른쪽/아래 글상자를 먼저 만든다: 읽는 순서는 위치 기준이어야 한다
    s.shapes.add_textbox(Inches(5), Inches(1), Inches(3), Inches(1)).text_frame.text = "오른쪽"
    s.shapes.add_textbox(Inches(1), Inches(4), Inches(3), Inches(1)).text_frame.text = "아래"
    s.shapes.add_textbox(Inches(1), Inches(1), Inches(3), Inches(1)).text_frame.text = "왼쪽"
    tbl = s.shapes.add_table(2, 2, Inches(1), Inches(6), Inches(4), Inches(1)).table
    tbl.cell(0, 0).text, tbl.cell(0, 1).text = "가", "나"
    tbl.cell(1, 0).text = "다"
    s._element.set("show", "0")

    path = tmp_path / "t.pptx"
    prs.save(path)
    return path


def test_load_deck(deck_path):
    deck = load_deck(deck_path)
    assert len(deck.slides) == 2
    first, second = deck.slides
    assert first.number == 1 and first.title == "제목"
    assert first.body == "제목\n첫 줄\n둘째 줄"
    assert first.notes == "노트 대본입니다."
    assert not first.hidden
    assert second.body == "왼쪽\n오른쪽\n아래\n가, 나\n다"
    assert second.notes == "" and second.hidden


def test_script_sources():
    s = Slide(1, title="T", body="T\n본문", notes="노트")
    assert s.script("notes_first") == "노트"
    assert s.script("body") == "T\n본문"
    assert s.script("notes") == "노트"
    assert s.script("both") == "T\n본문\n\n노트"
    assert Slide(2, body="본문만").script("notes_first") == "본문만"


def test_to_speech():
    assert to_speech("• 첫째\n- 둘째\n\n셋째입니다.\n질문?") == "첫째.\n둘째.\n셋째입니다.\n질문?"
    assert to_speech("  \n \v ") == ""
    assert to_speech("줄\v바꿈") == "줄.\n바꿈."


def test_resolve_local_path(tmp_path, monkeypatch):
    f = tmp_path / "Docs" / "발표.pptx"
    f.parent.mkdir()
    f.write_bytes(b"x")
    assert resolve_local_path(str(f)) == f
    assert resolve_local_path("Presentation1") is None
    monkeypatch.setenv("OneDrive", str(tmp_path))
    url = "https://d.docs.live.net/abc123/Docs/%EB%B0%9C%ED%91%9C.pptx"
    assert resolve_local_path(url) == f
