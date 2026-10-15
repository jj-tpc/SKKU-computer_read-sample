"""시연용 예제 발표 자료(sample/sample_presentation.pptx)를 만든다.

    python tools/make_sample.py
"""
from pathlib import Path

from pptx import Presentation
from pptx.util import Inches, Pt

OUT = Path(__file__).resolve().parent.parent / "sample" / "sample_presentation.pptx"


def add_notes(slide, text):
    slide.notes_slide.notes_text_frame.text = text


def main():
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)

    s = prs.slides.add_slide(prs.slide_layouts[0])
    s.shapes.title.text = "인공지능과 우리의 학교생활"
    s.placeholders[1].text = "컴퓨터 과목 발표 · 2학년 3반 홍길동"
    add_notes(s, "안녕하세요. 오늘은 인공지능이 우리의 학교생활을 어떻게 바꾸고 있는지 발표하겠습니다.")

    s = prs.slides.add_slide(prs.slide_layouts[1])
    s.shapes.title.text = "인공지능이란?"
    body = s.placeholders[1].text_frame
    body.text = "사람처럼 배우고 판단하는 컴퓨터 프로그램"
    for line in ("데이터에서 규칙을 스스로 찾아냄", "번역, 음성 인식, 그림 그리기 등에 활용"):
        body.add_paragraph().text = line
    add_notes(s, "인공지능은 사람처럼 배우고 판단하는 컴퓨터 프로그램입니다.\n"
                 "많은 데이터를 보고 스스로 규칙을 찾아낸다는 점이 일반 프로그램과 다릅니다.")

    # 발표자 노트가 없는 장: '노트 우선'이면 슬라이드 본문을 읽는다
    s = prs.slides.add_slide(prs.slide_layouts[5])
    s.shapes.title.text = "학교에서 쓰이는 예"
    rows = [("분야", "예시"), ("수업", "AI 튜터, 자동 채점"), ("생활", "급식 메뉴 추천, 번역기")]
    table = s.shapes.add_table(len(rows), 2, Inches(1), Inches(2), Inches(11), Inches(2.5)).table
    for r, row in enumerate(rows):
        for c, value in enumerate(row):
            table.cell(r, c).text = value

    # 숨긴 슬라이드: 자동 발표에서 건너뛴다
    s = prs.slides.add_slide(prs.slide_layouts[1])
    s.shapes.title.text = "(예비) 질문 대비 자료"
    s.placeholders[1].text = "인공지능이 틀릴 때는 어떻게 하나요?"
    s._element.set("show", "0")

    s = prs.slides.add_slide(prs.slide_layouts[6])  # 빈 레이아웃에 글상자 두 개
    left = s.shapes.add_textbox(Inches(1), Inches(1), Inches(5), Inches(2)).text_frame
    left.text = "좋은 점"
    left.add_paragraph().text = "• 궁금한 것을 바로 물어볼 수 있다"
    right = s.shapes.add_textbox(Inches(7), Inches(1), Inches(5), Inches(2)).text_frame
    right.text = "주의할 점"
    right.add_paragraph().text = "• 답이 항상 맞지는 않다"
    for tf in (left, right):
        tf.paragraphs[0].runs[0].font.size = Pt(28)
    add_notes(s, "인공지능은 편리하지만 답이 항상 맞는 것은 아닙니다. 꼭 다시 확인하는 습관이 필요합니다.")

    s = prs.slides.add_slide(prs.slide_layouts[0])
    s.shapes.title.text = "감사합니다"
    s.placeholders[1].text = "질문 있으신가요?"
    add_notes(s, "지금까지 발표를 들어 주셔서 감사합니다. 궁금한 점이 있으면 질문해 주세요.")

    OUT.parent.mkdir(exist_ok=True)
    prs.save(OUT)
    print("만들었습니다:", OUT)


if __name__ == "__main__":
    main()
