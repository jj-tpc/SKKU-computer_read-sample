"""PPTX 파일에서 슬라이드별 본문/발표자 노트를 뽑아내고, 음성으로 읽기 좋게 다듬는다."""
from __future__ import annotations

import re
from dataclasses import dataclass, field
from pathlib import Path

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE, PP_PLACEHOLDER

# 읽을 필요가 없는 자리표시자(바닥글, 날짜, 쪽 번호)
_SKIP_PLACEHOLDERS = {PP_PLACEHOLDER.FOOTER, PP_PLACEHOLDER.DATE, PP_PLACEHOLDER.SLIDE_NUMBER}

# 읽을 내용 선택지 (키, 화면 표시 이름)
SOURCES = [
    ("notes_first", "노트 우선 (없으면 본문)"),
    ("body", "슬라이드 본문"),
    ("notes", "발표자 노트"),
    ("both", "본문 + 노트"),
]


@dataclass
class Slide:
    number: int  # 1부터 시작하는 슬라이드 번호 (PowerPoint의 SlideIndex와 같음)
    title: str = ""
    body: str = ""  # 제목을 포함한 슬라이드 위 글자 전체
    notes: str = ""
    hidden: bool = False

    def script(self, source: str) -> str:
        """선택한 읽을 내용(source)에 맞는 대본을 돌려준다."""
        if source == "body":
            return self.body
        if source == "notes":
            return self.notes
        if source == "both":
            return "\n\n".join(t for t in (self.body, self.notes) if t)
        return self.notes or self.body  # notes_first


@dataclass
class Deck:
    path: Path
    slides: list[Slide] = field(default_factory=list)

    @property
    def name(self) -> str:
        return self.path.name


def _shape_texts(shape) -> list[str]:
    """도형 하나(그룹이면 그 안의 도형들까지)에서 글자를 꺼낸다."""
    if shape.is_placeholder and shape.placeholder_format.type in _SKIP_PLACEHOLDERS:
        return []
    if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
        out: list[str] = []
        for child in _in_reading_order(shape.shapes):
            out.extend(_shape_texts(child))
        return out
    if getattr(shape, "has_table", False) and shape.has_table:
        rows = []
        for row in shape.table.rows:
            cells = [_clean(c.text) for c in row.cells]
            cells = [c for c in cells if c]
            if cells:
                rows.append(", ".join(cells))
        return rows
    if shape.has_text_frame:
        text = _clean(shape.text_frame.text)
        return [text] if text else []
    return []


def _in_reading_order(shapes):
    """위에서 아래로, 같은 줄이면 왼쪽에서 오른쪽으로."""
    return sorted(shapes, key=lambda s: (s.top or 0, s.left or 0))


def _clean(text: str) -> str:
    # python-pptx는 줄바꿈(<a:br>)을 \v 로 돌려준다.
    lines = text.replace("\v", "\n").replace("\r", "\n").split("\n")
    return "\n".join(line.strip() for line in lines if line.strip())


def load_deck(path: str | Path) -> Deck:
    path = Path(path)
    prs = Presentation(str(path))
    deck = Deck(path=path)
    for number, s in enumerate(prs.slides, start=1):
        title_shape = s.shapes.title
        title = _clean(title_shape.text_frame.text) if title_shape is not None else ""
        parts = [title] if title else []
        for shape in _in_reading_order(s.shapes):
            if title_shape is not None and shape.shape_id == title_shape.shape_id:
                continue
            parts.extend(_shape_texts(shape))
        notes = ""
        if s.has_notes_slide and s.notes_slide.notes_text_frame is not None:
            notes = _clean(s.notes_slide.notes_text_frame.text)
        deck.slides.append(Slide(
            number=number,
            title=title,
            body="\n".join(parts),
            notes=notes,
            hidden=s._element.get("show") == "0",
        ))
    return deck


_BULLET = re.compile(r"^[\s•·∙◦▪▫■□◆◇○●◎▶▷►➢➤→⇒※✓✔☞*\-–—]+")
_SENTENCE_END = tuple(".?!。…:;,)]}\"'”’")


def to_speech(text: str) -> str:
    """글머리 기호를 지우고, 줄 끝마다 마침표를 붙여 TTS가 줄 사이에서 쉬도록 만든다."""
    lines = []
    for line in text.replace("\v", "\n").replace("\r", "\n").split("\n"):
        line = _BULLET.sub("", line).strip()
        line = re.sub(r"\s+", " ", line)
        if not line:
            continue
        if not line.endswith(_SENTENCE_END):
            line += "."
        lines.append(line)
    return "\n".join(lines)
