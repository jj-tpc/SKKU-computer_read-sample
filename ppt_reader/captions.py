"""캡션: 슬라이드 속 그림을 꺼내 LLM에게 설명을 받고, 그림 해시로 저장한다. 질의응답 근거에만 쓴다.

그림 꺼내기는 python-pptx + Pillow로 한다(새 의존성 없음). 근거: docs/research/slide-images.md
"""
from __future__ import annotations

import hashlib
import io
import json
import queue
import re
import sys
import threading
import time
from collections import Counter
from dataclasses import dataclass, field
from pathlib import Path
from typing import Callable

from .llm.base import LLMError, Request
from .qa import CAPTION_SYSTEM
from .settings import _write_atomic

LONG_EDGE = 1024
MAX_PICTURES = 40
TINY_EDGE = 64  # 원본 긴 변이 이보다 작으면 아주 작은 그림
TINY_AREA = 0.01  # 슬라이드 위 넓이가 슬라이드의 1% 미만이면 아주 작은 그림
REPEAT_LIMIT = 3  # 같은 그림이 이보다 많은 장에 나오면 로고·장식으로 본다
METAFILES = {"image/x-emf": "EMF", "image/x-wmf": "WMF"}
_NS = {"a": "http://schemas.openxmlformats.org/drawingml/2006/main",
       "p": "http://schemas.openxmlformats.org/presentationml/2006/main",
       "r": "http://schemas.openxmlformats.org/officeDocument/2006/relationships",
       "asvg": "http://schemas.microsoft.com/office/drawing/2016/SVG/main"}

SKIP_SVG = "그림을 꺼낼 수 없는 형식(SVG)이라 건너뜀"
SKIP_LINK = "연결된 그림이라 건너뜀"
SKIP_META = "이 컴퓨터에서 바꿀 수 없는 형식({fmt})이라 건너뜀"
SKIP_TINY = "아주 작은 그림이라 건너뜀"
SKIP_REPEAT = "여러 장에 되풀이되는 그림(로고·장식)이라 건너뜀"
OVER_LIMIT = "앞 40개만 만들었어요"


@dataclass
class SlidePicture:
    slide: int  # 슬라이드 번호(1부터)
    index: int  # 그 장 안의 그림 차례(1부터)
    key: str = ""  # sha256(원본 바이트) + 잘라내기 값
    blob: bytes = b""
    content_type: str = ""
    crop: tuple[float, float, float, float] = (0.0, 0.0, 0.0, 0.0)
    size: tuple[int, int] = (0, 0)
    area: float = 1.0  # 슬라이드 넓이에 대한 비율
    skip: str = ""  # 건너뛰는 까닭(대본 창에 보인다)


def picture_key(blob: bytes, crop) -> str:
    return hashlib.sha256(blob).hexdigest() + ":" + ",".join(f"{max(0.0, c):.4f}" for c in crop)


def _image_size(blob: bytes) -> tuple[int, int]:
    from PIL import Image

    try:
        return Image.open(io.BytesIO(blob)).size
    except Exception:
        return (0, 0)


def _from_part(slide_no: int, part, crop=(0.0, 0.0, 0.0, 0.0), area: float = 1.0) -> SlidePicture:
    pic = SlidePicture(slide_no, 0, blob=part.blob, content_type=part.content_type, crop=tuple(crop), area=area)
    pic.key = picture_key(pic.blob, pic.crop)
    pic.size = _image_size(pic.blob)
    return pic


def extract(prs) -> dict[int, list[SlidePicture]]:
    """장마다 그림 목록. 그룹 안, 그림 자리표시자, 배경 그림, 그림으로 채운 도형까지 꺼낸다."""
    from lxml import etree
    from pptx.enum.shapes import MSO_SHAPE_TYPE
    from pptx.shapes.picture import Picture

    slide_area = float(prs.slide_width * prs.slide_height) or 1.0
    out: dict[int, list[SlidePicture]] = {}
    for number, slide in enumerate(prs.slides, 1):
        pics: list[SlidePicture] = []
        for rid in slide._element.xpath("./p:cSld/p:bg/p:bgPr/a:blipFill/a:blip/@r:embed"):
            pics.append(_from_part(number, slide.part.related_part(rid)))

        def walk(shapes):
            for sh in shapes:
                if sh.shape_type == MSO_SHAPE_TYPE.GROUP:
                    walk(sh.shapes)
                    continue
                area = (sh.width or 0) * (sh.height or 0) / slide_area
                if isinstance(sh, Picture):  # shape_type == PICTURE로 거르면 그림 자리표시자를 놓친다
                    rid = sh._element.blip_rId
                    if rid is None:
                        # python-pptx의 xpath는 자기 네임스페이스 표만 써서 asvg를 모른다. lxml의 xpath를 직접 부른다
                        svg = etree.ElementBase.xpath(sh._element, ".//asvg:svgBlip", namespaces=_NS)
                        pics.append(SlidePicture(number, 0, area=area, skip=SKIP_SVG if svg else SKIP_LINK))
                        continue
                    crop = (sh.crop_left, sh.crop_top, sh.crop_right, sh.crop_bottom)
                    # 형식은 그림 부분의 content_type으로 본다(python-pptx는 EMF를 WMF로 잘못 알려 준다)
                    pics.append(_from_part(number, sh.part.related_part(rid), crop, area))
                    continue
                for rid in sh._element.xpath("./p:spPr/a:blipFill/a:blip/@r:embed"):
                    pics.append(_from_part(number, sh.part.related_part(rid), area=area))

        walk(slide.shapes)
        for i, pic in enumerate(pics, 1):
            pic.index = i
        out[number] = pics
    _mark_skips(out)
    return out


def _mark_skips(slides: dict[int, list[SlidePicture]]) -> None:
    seen = Counter()
    for pics in slides.values():
        for key in {p.key for p in pics if p.key}:
            seen[key] += 1
    for pics in slides.values():
        for p in pics:
            if p.skip:
                continue
            if p.content_type in METAFILES and sys.platform != "win32":  # Pillow는 Windows에서만 그린다
                p.skip = SKIP_META.format(fmt=METAFILES[p.content_type])
            elif p.content_type == "image/svg+xml":
                p.skip = SKIP_SVG
            elif (p.size[0] and max(p.size) < TINY_EDGE) or p.area < TINY_AREA:
                p.skip = SKIP_TINY
            elif seen[p.key] > REPEAT_LIMIT:
                p.skip = SKIP_REPEAT


def to_llm_image(pic: SlidePicture) -> bytes:
    """잘라내기를 적용하고 긴 변 1024px 이하로 줄여 PNG(투명이 있으면)나 JPEG로."""
    from PIL import Image

    im = Image.open(io.BytesIO(pic.blob))
    if pic.content_type in METAFILES:
        dpi = im.info.get("dpi", 72)
        dpi = dpi[0] if isinstance(dpi, tuple) else dpi
        im.load(dpi=max(1, int(dpi * LONG_EDGE / max(im.size))))
    l, t, r, b = (max(0.0, c) for c in pic.crop)  # 음수 잘라내기(그림 밖으로 넓힘)는 0으로 본다
    if any((l, t, r, b)):
        w, h = im.size
        im = im.crop((round(w * l), round(h * t), round(w * (1 - r)), round(h * (1 - b))))
    if max(im.size) > LONG_EDGE:
        im.thumbnail((LONG_EDGE, LONG_EDGE), Image.Resampling.LANCZOS)
    out = io.BytesIO()
    has_alpha = im.mode in ("RGBA", "LA") or (im.mode == "P" and "transparency" in im.info)
    if has_alpha:
        im.convert("RGBA").save(out, "PNG")
    else:
        im.convert("RGB").save(out, "JPEG", quality=85)
    return out.getvalue()


def plan(slides: dict[int, list[SlidePicture]]) -> tuple[list[SlidePicture], bool]:
    """캡션을 만들 그림(슬라이드 차례로 앞 40개)과, 40개가 넘었는지."""
    todo = [p for n in sorted(slides) for p in slides[n] if not p.skip]
    return todo[:MAX_PICTURES], len(todo) > MAX_PICTURES


def caption_request(model: str, slide_text: str, pictures: list[SlidePicture]) -> Request:
    question = f"슬라이드의 글:\n{slide_text or '(없음)'}\n\n그림 {len(pictures)}개를 차례대로 설명해 줘."
    return Request(model, "low", CAPTION_SYSTEM, "", [], question, tuple(to_llm_image(p) for p in pictures))


_LINE = re.compile(r"^\s*그림\s*(\d+)\s*[:：]\s*(.+?)\s*$")


def parse_captions(text: str, count: int) -> dict[int, str]:
    """"그림 k: ..." 줄을 그림 차례마다 읽는다."""
    out: dict[int, str] = {}
    for line in text.splitlines():
        m = _LINE.match(line)
        if m and 1 <= int(m.group(1)) <= count:
            out[int(m.group(1))] = m.group(2)
    return out


class CaptionStore:
    """captions.json: {그림 키: {"text", "edited", "model", "made_at"}}. 파일 이름을 바꾸거나 옮겨도 다시 쓴다."""

    def __init__(self, folder: Path | str):
        self.path = Path(folder) / "captions.json"
        try:
            data = json.loads(self.path.read_text(encoding="utf-8"))
        except (OSError, ValueError):
            data = {}
        self.data: dict = data if isinstance(data, dict) else {}
        self._lock = threading.Lock()

    def get(self, key: str) -> dict | None:
        e = self.data.get(key)
        return e if isinstance(e, dict) and isinstance(e.get("text"), str) else None

    def put(self, key: str, text: str, model: str, edited: bool = False) -> None:
        with self._lock:
            self.data[key] = {"text": text, "edited": edited, "model": model,
                              "made_at": time.strftime("%Y-%m-%dT%H:%M:%S")}
            _write_atomic(self.path, self.data)

    def edit(self, key: str, text: str) -> None:
        """대본 창에서 고친 캡션. [다시 만들기]를 해도 덮어쓰지 않는다."""
        old = self.get(key) or {}
        self.put(key, text, old.get("model", ""), edited=True)

    def texts(self, slides: dict[int, list[SlidePicture]]) -> dict[int, list[str]]:
        """근거 묶음용: 장마다 그림 차례대로 캡션(없으면 빈 글)."""
        return {n: [(self.get(p.key) or {}).get("text", "") if p.key else "" for p in pics]
                for n, pics in slides.items()}


@dataclass
class _Job:
    model: str
    texts: dict[int, str]  # 장 번호 -> 슬라이드 글
    todo: list[SlidePicture] = field(default_factory=list)
    redo: bool = False  # [다시 만들기]: 고치지 않은 캡션을 다시 만든다


class CaptionWorker:
    """뒤에서 차례로 도는 캡션 작업 스레드 하나. 장마다 그 장의 그림을 한 번에 보낸다.

    on_progress(done, total)와 on_failed(LLMError)는 이 스레드에서 불린다. 화면 쪽이 화면 스레드로 옮긴다.
    """

    def __init__(self, store: CaptionStore, provider_factory: Callable, on_progress: Callable[[int, int], None],
                 on_failed: Callable[[LLMError], None]):
        self.store = store
        self.provider_factory = provider_factory  # () -> Provider. 지금 설정의 제공자
        self.on_progress = on_progress
        self.on_failed = on_failed
        self._cmds: queue.Queue = queue.Queue()
        self._cancel = threading.Event()
        self._thread: threading.Thread | None = None

    def start(self, model: str, texts: dict[int, str], todo: list[SlidePicture], redo: bool = False) -> None:
        self.stop()
        if self._thread is None:
            self._thread = threading.Thread(target=self._run, name="captions", daemon=True)
            self._thread.start()
        self._cmds.put(_Job(model, texts, list(todo), redo))

    def stop(self) -> None:
        self._cancel.set()

    def close(self) -> None:
        self.stop()
        if self._thread is not None:
            self._cmds.put(None)
            self._thread.join(timeout=3)

    def _run(self) -> None:
        while True:
            job = self._cmds.get()
            if job is None:
                return
            while not self._cmds.empty():  # 새로 온 일이 있으면 그것만 한다
                job = self._cmds.get()
                if job is None:
                    return
            self._cancel.clear()
            self._do(job)

    def _needs(self, pic: SlidePicture, redo: bool) -> bool:
        e = self.store.get(pic.key)
        return e is None or (redo and not e.get("edited"))

    def _do(self, job: _Job) -> None:
        total = len(job.todo)
        done = sum(1 for p in job.todo if not self._needs(p, job.redo))
        self.on_progress(done, total)
        by_slide: dict[int, list[SlidePicture]] = {}
        for p in job.todo:
            if self._needs(p, job.redo):
                by_slide.setdefault(p.slide, []).append(p)
        if not by_slide:
            return
        try:
            provider = self.provider_factory()
        except Exception as e:
            self.on_failed(LLMError("other", str(e)[:200]))
            return
        for number in sorted(by_slide):
            if self._cancel.is_set():
                return
            pics = by_slide[number]
            try:
                req = caption_request(job.model, job.texts.get(number, ""), pics)
                answer = "".join(provider.stream(req, self._cancel))
            except LLMError as e:
                self.on_failed(e)
                return
            except Exception as e:
                self.on_failed(LLMError("other", str(e)[:200]))
                return
            if self._cancel.is_set():
                return
            got = parse_captions(answer, len(pics))
            for k, pic in enumerate(pics, 1):
                if k in got and not (self.store.get(pic.key) or {}).get("edited"):
                    self.store.put(pic.key, got[k], job.model)
                    done += 1
            self.on_progress(done, total)
