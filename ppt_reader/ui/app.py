"""QApplication 시작, session과 화면 신호 잇기, --self-check. 화면 쪽 로직은 두지 않는다."""
from __future__ import annotations

import os
import sys
import threading
import time
from pathlib import Path
from typing import Callable

from PySide6.QtCore import QObject, QRect, QTimer, QUrl, qVersion
from PySide6.QtGui import QCursor, QDesktopServices, QGuiApplication, QScreen
from PySide6.QtWidgets import QApplication, QFileDialog, QMessageBox

from ..captions import CaptionStore, CaptionWorker, extract, plan

from ..follow import Area, Follower, wait_spot
from ..llm.base import LLMError, make_provider, start_stream
from ..session import CHIP_OFF, READING, SPEAKING, THINKING, Session
from ..settings import DOGS, Settings, app_folder
from ..slides import to_speech
from ..voice.cache import sentence_plan
from ..voice.cloud import ElevenLabsEngine, VoiceError, cloud_engine
from . import platform_mac, platform_win, theme
from .bubble import BubbleWindow
from .dog import FRAMES, SWEAT_ROW, DogWindow, load_sheet
from .orb import OrbWindow, QuestionWindow, corner_center
from .qtcore import MainThreadInvoker, QtClock

APP_TITLE = "PPT 읽어주기"
FRAME_MS = 33  # 강아지와 원 버튼을 함께 그리는 타이머(약 30fps)
HAPPY_SECONDS = 1.5
MONITOR_CHECK_MS = 1000  # "쇼가 뜬 모니터 따라가기"일 때 쇼 창의 모니터를 다시 보는 간격
PREVIEW_TEXT = "안녕하세요. 이 목소리로 읽어 드릴게요."
DEFAULT_OUTPUT = "시스템 기본 장치"
FILE_DIALOG_TITLE = "발표할 PPT 파일 선택"
FILE_FILTER = "PowerPoint 프레젠테이션 (*.pptx *.pptm);;모든 파일 (*.*)"


def screen_for_native_rect(rect: tuple[int, int, int, int]) -> QScreen | None:
    """Windows 모니터 사각형(실제 픽셀)에 맞는 Qt 화면. Qt는 화면 위치를 실제 픽셀 그대로 두고 크기만 배율로 나눈다."""
    left, top, right, bottom = rect
    cx, cy = (left + right) // 2, (top + bottom) // 2
    for screen in QGuiApplication.screens():
        g, dpr = screen.geometry(), screen.devicePixelRatio()
        native = QRect(g.x(), g.y(), round(g.width() * dpr), round(g.height() * dpr))
        if (g.x(), g.y()) == (left, top) or native.contains(cx, cy):
            return screen
    return None


TITLE_BAR = 40  # 창 제목 막대 몫(논리 px). 줄일 때 빼 둔다


def fit_on_screen(window, first: bool) -> None:
    """창이 화면 작업 영역보다 크면 줄인다(내용은 스크롤된다). 처음 열 때는 커서가 있는 화면 가운데에 둔다.

    화면 배율 150%인 1080p 모니터는 작업 영역이 약 1280×688(논리 px)이라, 줄이지 않으면 설정 창 아래가 작업 표시줄에 걸린다.
    """
    screen = QGuiApplication.screenAt(QCursor.pos()) or QGuiApplication.primaryScreen()
    if screen is None:
        return
    avail = screen.availableGeometry()
    w = min(window.width(), avail.width() - 16)
    h = min(window.height(), avail.height() - TITLE_BAR - 8)
    window.resize(w, h)
    if first:
        window.move(avail.x() + (avail.width() - w) // 2, avail.y() + max(0, (avail.height() - h - TITLE_BAR) // 2))


class Controller(QObject):
    """창들을 띄우고 session의 알림을 화면에 옮긴다."""

    def __init__(self, settings: Settings, session: Session, pipeline=None, presynth=None,
                 closers: list[Callable[[], None]] | None = None):
        super().__init__()
        self.settings = settings
        self.session = session
        self.pipeline = pipeline
        self.presynth = presynth
        self.presynth_state: tuple[str, int, int, str] = ("off", 0, 0, "")  # (상태, 만든 수, 전체, 실패 이유). 설정 창이 보인다
        self.closers = closers or []
        self.dog = DogWindow(settings["dog"])
        self.orb = OrbWindow(self.on_action)
        self.orb.set_speed(settings["speed"])
        self.question = QuestionWindow(self._send_question, self._question_closed)
        app = QGuiApplication.instance()
        if app is not None:  # 진단: 이 앱이 언제 활성이 되는지(PPT_READER_TRACE)
            app.applicationStateChanged.connect(lambda state: platform_mac.trace("app state", state.name))
        self.bubble = BubbleWindow()
        self._bubble_text = ""
        self.settings_window = None
        self.script_window = None
        self.invoker: MainThreadInvoker | None = None  # run()이 넣는다. 작업 스레드의 결과를 화면 스레드로
        self.link_chip = CHIP_OFF  # 연동 중일 때 session이 알려 준 마지막 칩 문구
        self.pictures: dict = {}  # 장 번호 -> 그림 목록(캡션)
        self.caption_store = CaptionStore(settings.folder)
        self.caption_worker: CaptionWorker | None = None
        self.caption_state: tuple[str, int, int, str] = ("off", 0, 0, "")
        self.caption_over_limit = False
        self._el_voices: list = []
        self._el_loading = False
        self.follower = Follower()
        self._screen_now: QScreen | None = None
        self._area = Area(0, 0, 1, 1)
        self._happy_until = 0.0
        self._last = time.monotonic()
        self._timer = QTimer(self)
        self._timer.setInterval(FRAME_MS)
        self._timer.timeout.connect(self._frame)
        self._monitor_timer = QTimer(self)
        self._monitor_timer.setInterval(MONITOR_CHECK_MS)
        self._monitor_timer.timeout.connect(self._place)
        session.on_state = self._on_state
        session.on_pose = self._on_pose
        session.on_message = self._on_message
        session.on_notice = self.orb.notify
        session.on_shake = self.orb.shake
        session.on_bubble = self._on_bubble
        session.on_question = self._on_question
        if pipeline is not None:
            pipeline.on_started = session.voice_started
        session.on_deck = self._on_deck
        session.on_link = self._on_link
        session.on_slide = self._on_slide
        session.on_script = self._on_script

    def start(self) -> None:
        self._place()
        orb = self.orb.orb_center_global()
        self.follower.pos = wait_spot((orb.x(), orb.y()), self._area)
        self.dog.move(*map(round, self.follower.pos))
        if self.settings["dog_visible"]:
            self.dog.show()
        self.orb.show()
        self.orb.raise_()  # 원 버튼 묶음은 늘 강아지보다 위
        self._timer.start()
        self._monitor_timer.start()
        if self.settings.first_run:  # 처음 켰을 때: 강아지가 기뻐하고 설정 창이 저절로 열린다
            self._on_pose("happy")
            self.open_settings(first_run=True)
            self.settings.set(first_run_done=True)

    def _screen(self) -> QScreen:
        """설정의 원 버튼 모니터. 쇼 따라가기면 쇼 창이 있는 모니터, 모르면 주 모니터(macOS는 늘 주 모니터)."""
        name = self.settings["monitor"]
        if name == "follow":
            rect = platform_win.show_monitor_rect()
            found = screen_for_native_rect(rect) if rect else None
            return (found or (self._screen_now if self._screen_now in QGuiApplication.screens() else None)
                    or QGuiApplication.primaryScreen())
        for screen in QGuiApplication.screens():
            if screen.name() == name:
                return screen
        return QGuiApplication.primaryScreen()

    def _place(self) -> None:
        platform_mac.note_front()
        screen = self._screen()
        available = screen.availableGeometry()
        area = Area(available.x(), available.y(), available.width(), available.height())
        if screen is self._screen_now and area == self._area:
            return
        self._screen_now, self._area = screen, area
        self.orb.place(available)

    def _frame(self) -> None:
        now = time.monotonic()
        dt, self._last = min(0.1, now - self._last), now
        cursor = QCursor.pos()
        frame = self.pipeline.frame() if self.pipeline is not None else None
        speaking = frame is not None and frame.speaking
        state = self.session.state
        orb = self.orb.orb_center_global()
        over = self.orb.is_over(cursor) or (self.question.isVisible() and self.question.geometry().contains(cursor))
        # 질의응답 중과 말풍선이 떠 있는 동안은 그 자리에 머문다. 말풍선이 사라지면 다시 커서에게 간다
        stay = state in (THINKING, SPEAKING) or self.session.bubble
        base = self.follower.update(now, dt, (cursor.x(), cursor.y()), self._area, (orb.x(), orb.y()), over, stay)
        if now < self._happy_until:
            pose = "happy"
        elif state == THINKING:
            pose = "think"
        elif speaking and not (state == READING and self.follower.moving):
            pose = "talk"  # 대본을 읽는 동안에는 따라가며 걷고, 멈추면 말하는 입을 한다
        else:
            pose = base
        self.dog.set_pose(pose)
        if speaking:
            self.dog.mouth = frame.mouth
        self.dog.sweat = self.session.sweat
        self.dog.facing_left = self.follower.facing_left
        self.dog.move(*map(round, self.follower.pos))
        if self.session.bubble:
            if speaking:  # 목소리가 읽기 시작한 낱말까지, 앞 문장들에 이어서
                said = " ".join(frame.text.split()[:frame.words])
                self._bubble_text = f"{frame.before} {said}".strip()
            self.bubble.set_text(self._bubble_text)
            self.bubble.place(*self.follower.pos, self._area)
        self.bubble.tick(dt)
        self.question.tick(dt)
        level, tone = (frame.level, frame.tone) if speaking else (0.0, 0.5)
        self.orb.level += (level - self.orb.level) * 0.35
        self.orb.tone += (tone - self.orb.tone) * 0.05  # 높낮이 느낌은 천천히 따라간다
        self.dog.tick(dt)
        self.orb.tick(dt, cursor)

    def _on_state(self, state: str) -> None:
        self.orb.set_state(state)
        engine = getattr(self.pipeline, "os_engine", None)
        factors = dict(getattr(engine, "x5_factors", None) or {})
        if state == "idle" and factors and factors != self.settings["x5_factor"]:
            self.settings.set(x5_factor=factors)  # 음성마다 처음 한 번 잰 x5 배속을 기억한다

    # ------------------------------------------------------------ 대본 소리 미리 만들기

    def start_presynth(self) -> None:
        if self.presynth is None:
            return
        plan = cloud_engine(self.settings.data)
        deck = self.session.deck
        if plan is None or deck is None:
            self.presynth.stop()
            self.presynth_state = ("off", 0, 0, "")
            return
        slides = [sentence_plan(to_speech(self.session.script_for(i))) for i in range(len(deck.slides))]
        self.presynth.start(plan[0], plan[1], slides, max(self.session.current, 0))

    def _on_script(self, index: int) -> None:
        if index < 0 or self.presynth is None or cloud_engine(self.settings.data) is None:
            self.start_presynth()
        else:
            self.presynth.redo(index, sentence_plan(to_speech(self.session.script_for(index))))

    def presynth_progress(self, done: int, total: int) -> None:
        self.presynth_state = ("done" if done >= total else "making", done, total, "")
        self._refresh_windows()

    def presynth_failed(self, kind: str, detail: str) -> None:
        _, done, total, _ = self.presynth_state
        self.presynth_state = ("failed", done, total, detail or kind)
        self._refresh_windows()

    def _on_pose(self, pose: str) -> None:
        if pose == "happy":
            self._happy_until = time.monotonic() + HAPPY_SECONDS

    def _on_message(self, kind: str, text: str) -> None:
        if kind == "error":
            QMessageBox.critical(None, APP_TITLE, text)
        elif kind == "warning":
            QMessageBox.warning(None, APP_TITLE, text)

    # ------------------------------------------------------------ 질의응답

    def open_question(self) -> None:
        if not self.question.isVisible():
            self._question_from = platform_win.foreground_window()  # 닫을 때 쇼에 포커스를 돌려준다
            self._question_from_mac = platform_mac.front_app()
        self.question.place(self.orb.orb_center_global())
        self.orb.question_open = True
        self.question.open()

    def _send_question(self, text: str, effort: str) -> None:
        self.session.ask(text, effort)

    def _question_closed(self) -> None:
        self.orb.question_open = False
        platform_win.restore_foreground(getattr(self, "_question_from", 0))
        platform_mac.restore_front(getattr(self, "_question_from_mac", None))

    def _on_question(self, text: str) -> None:
        self.question.set_text(text)  # ""면 비우고(답을 다 함), 글이면 다시 채운다(오류)

    def _on_bubble(self, visible: bool) -> None:
        if visible:
            self._bubble_text = ""
            self.bubble.begin()
            self.bubble.place(*self.follower.pos, self._area)
        else:
            self.bubble.hide()

    # ------------------------------------------------------------ 설정 창, 대본 창

    def _post(self, fn, *args) -> None:
        if self.invoker is not None:
            self.invoker.post(fn, args)
        else:
            fn(*args)

    def _refresh_windows(self) -> None:
        if self.settings_window is not None and self.settings_window.isVisible():
            self.settings_window.refresh()
        if self.script_window is not None and self.script_window.isVisible():
            self.script_window.show_captions()
            self.script_window.refresh_cards()

    def open_settings(self, first_run: bool = False) -> None:
        from .settings_window import SettingsWindow

        created = self.settings_window is None
        if created:
            self.settings_window = SettingsWindow(self)
        self.settings_window.refresh()
        self.settings_window.open_cards(first_run)
        fit_on_screen(self.settings_window, created)
        self.settings_window.show()
        self.settings_window.raise_()
        self.settings_window.activateWindow()

    def open_script_window(self) -> None:
        from .script_window import ScriptWindow

        if self.session.deck is None:
            return
        created = self.script_window is None
        if created:
            self.script_window = ScriptWindow(self)
        else:
            self.script_window.rebuild()
        fit_on_screen(self.script_window, created)
        self.script_window.show()
        self.script_window.raise_()
        self.script_window.activateWindow()

    def open_file_dialog(self) -> None:
        path, _ = QFileDialog.getOpenFileName(self.settings_window, FILE_DIALOG_TITLE, "", FILE_FILTER)
        if path:
            self.session.open_file(path)

    def open_licenses(self) -> None:
        """licenses 폴더를 연다. 실행 파일 안(PyInstaller가 푼 곳)과 저장소 루트 순서로 찾는다."""
        roots = [Path(getattr(sys, "_MEIPASS", "")), Path(sys.executable).parent, Path(__file__).resolve().parents[2]]
        for base in roots:
            folder = base / "licenses"
            if str(base) and folder.is_dir():
                QDesktopServices.openUrl(QUrl.fromLocalFile(str(folder)))
                return

    def apply(self, **values) -> None:
        """설정 창에서 바꾼 값을 바로 저장하고 뒤따르는 일을 한다."""
        self.session.set_options(**values)
        if "dog" in values:
            self.dog.set_dog(values["dog"])
        if "dog_visible" in values:
            self.dog.setVisible(values["dog_visible"])
        if "monitor" in values:
            self._screen_now = None
            self._place()
        if {"voice_engine", "cloud_voice"} & values.keys():
            self.start_presynth()
        if {"captions", "provider", "models"} & values.keys():
            self.start_captions()

    def set_link(self, on: bool) -> None:
        self.session.set_link(on)
        self._refresh_windows()

    def set_key(self, owner: str, key: str) -> None:
        self.settings.set_for("keys", owner, key)
        self.settings.set_for("key_ok", owner, False)  # 새 키는 확인하기 전까지 쓰지 않는다
        self._after_key(owner)

    def check_key(self, owner: str, key: str, done: Callable[[bool | None, str], None]) -> None:
        """그 제공자의 모델(목소리) 목록을 한 번 불러 본다. 맞음 True, 키 문제 False, 그 밖 None."""
        def work():
            try:
                if owner == "elevenlabs":
                    ElevenLabsEngine(key, "", "").check_key()
                else:
                    make_provider(owner, key).check_key()
                result = (True, "")
            except (LLMError, VoiceError) as e:
                result = (False, e.detail) if e.kind == "key" else (None, e.detail or e.kind)
            except Exception as e:
                result = (None, str(e)[:120])
            self._post(self._key_checked, owner, key, result, done)

        threading.Thread(target=work, name="key-check", daemon=True).start()

    def _key_checked(self, owner, key, result, done) -> None:
        ok, reason = result
        if key == self.settings["keys"].get(owner):
            self.settings.set_for("key_ok", owner, ok is True)
            self._after_key(owner)
        done(ok, reason)

    def _after_key(self, owner: str) -> None:
        if owner == "elevenlabs":
            self._el_voices = []
        self.start_presynth()
        self.start_captions()
        self._refresh_windows()

    def preview_voice(self) -> None:
        if self.pipeline is None:
            return
        self.session.stop()
        self.pipeline.speak(PREVIEW_TEXT, self.session.voice_id(), 1.0, tag=("preview",),
                            output_id=self.session.output_id())

    def os_voices(self) -> list[tuple[str, str, str]]:
        voices = sorted(self.session.voice_list, key=lambda v: not v.is_korean)  # 한국어 음성 먼저
        return [(v.id, f"{v.name}  [{v.lang}]" if v.lang else v.name, v.name) for v in voices]

    def elevenlabs_voices(self) -> list:
        """계정의 목소리 목록. 처음 부르면 뒤에서 불러오고, 다 오면 설정 창을 다시 그린다."""
        s = self.settings
        if not self._el_voices and not self._el_loading and s["key_ok"].get("elevenlabs") and s["keys"].get("elevenlabs"):
            self._el_loading = True
            key = s["keys"]["elevenlabs"]

            def work():
                try:
                    voices = ElevenLabsEngine(key, "", "").voices()
                except VoiceError:
                    voices = []
                self._post(self._el_loaded, voices)

            threading.Thread(target=work, name="el-voices", daemon=True).start()
        return list(self._el_voices)

    def _el_loaded(self, voices) -> None:
        self._el_voices, self._el_loading = voices, False
        if voices and not self.settings["cloud_voice"]["elevenlabs"]["voice_id"]:
            self.session.cloud_voice_chosen("elevenlabs", voices[0])  # 처음에는 첫 목소리
        self._refresh_windows()

    def output_choices(self):
        from ..voice.player import AudioOutput, default_output, list_outputs

        default = default_output()
        first = DEFAULT_OUTPUT + (f" (지금: {default.name})" if default else "")
        out, seen = [], set()
        for o in list_outputs():
            text, n = o.name, 2
            while text in seen:  # 이름이 같은 장치가 여러 개면 번호를 붙인다
                text, n = f"{o.name} ({n})", n + 1
            seen.add(text)
            out.append((o, text))
        saved = self.settings["output"]
        chosen = AudioOutput(saved["id"], saved["name"]) if saved else None
        if chosen is not None and not any(o.id.lower() in chosen.id.lower() or chosen.id.lower() in o.id.lower()
                                          for o, _ in out):
            out.append((chosen, f"{chosen.name} (연결 안 됨 - 기본 장치로 읽음)"))
        elif chosen is not None:
            chosen = next(o for o, _ in out if o.id.lower() in chosen.id.lower() or chosen.id.lower() in o.id.lower())
        return first, out, chosen

    def monitor_choices(self) -> list[tuple[str, str]]:
        items = [("follow", "쇼가 뜬 모니터 따라가기")]
        for i, screen in enumerate(QGuiApplication.screens(), 1):
            size, dpr = screen.size(), screen.devicePixelRatio()
            items.append((screen.name(), f"{i}번 — {screen.name()} ({round(size.width() * dpr)}×{round(size.height() * dpr)})"))
        return items

    def presynth_text(self) -> str:
        state, done, total, reason = self.presynth_state
        if state == "making":
            return f"대본 소리 만드는 중 {done} / {total}"
        if state == "done":
            return "대본 소리 준비됨"
        if state == "failed":
            return f"대본 소리를 만들지 못했어요: {reason}"
        return ""

    # ------------------------------------------------------------ 파일, 연동, 캡션

    def _on_deck(self, deck) -> None:
        self.start_presynth()
        self.pictures = {}
        self.session.captions = {}
        self.caption_over_limit = False
        if self.script_window is not None and self.script_window.isVisible():
            self.script_window.rebuild()
        self._refresh_windows()
        if deck is None:
            return
        path = deck.path

        def work():  # 큰 파일은 그림을 꺼내는 데 시간이 걸려서 뒤에서 한다
            try:
                from pptx import Presentation

                pictures = extract(Presentation(path))
            except Exception:
                pictures = {}
            self._post(self._pictures_ready, path, pictures)

        threading.Thread(target=work, name="pictures", daemon=True).start()

    def _pictures_ready(self, path, pictures) -> None:
        deck = self.session.deck
        if deck is None or deck.path != path:
            return  # 그 사이 다른 파일을 열었다
        self.pictures = pictures
        self.session.captions = self.caption_store.texts(pictures)
        if self.script_window is not None and self.script_window.isVisible():
            self.script_window.rebuild()
        self.start_captions()

    def _on_link(self, text: str) -> None:
        self.link_chip = text
        self._refresh_windows()

    def link_chip_state(self) -> tuple[str, str]:
        """설정 창의 연동 칩: (글, 색). 연동이 꺼져 있으면 파일이 있나에 따라 "연동 꺼짐" 또는 "파일 없음"."""
        if self.session.link is None:
            return (CHIP_OFF if self.session.deck is not None else "파일 없음"), "off"
        text = self.link_chip
        return text, ("ok" if text.startswith("PowerPoint 연결됨") else "bad" if "끊겼" in text else "warn")

    def _on_slide(self, index: int) -> None:
        if self.script_window is not None and self.script_window.isVisible():
            self.script_window.follow(index)

    def start_captions(self, redo: bool = False) -> None:
        s = self.settings
        if self.caption_worker is not None:
            self.caption_worker.stop()
        deck = self.session.deck
        if not s["captions"]:
            self.caption_state = ("off", 0, 0, "")
        elif not s["key_ok"].get(s["provider"]):
            self.caption_state = ("paused", 0, 0, "")
        elif deck is not None and self.pictures and self.caption_worker is not None:
            todo, self.caption_over_limit = plan(self.pictures)
            texts = {sl.number: sl.body for sl in deck.slides}
            self.caption_state = ("making", 0, len(todo), "")
            self.caption_worker.start(s["models"][s["provider"]], texts, todo, redo)
        self._refresh_windows()

    def stop_captions(self) -> None:
        if self.caption_worker is not None:
            self.caption_worker.stop()
        _, done, total, _ = self.caption_state
        self.caption_state = ("stopped", done, total, "")
        self._refresh_windows()

    def redo_captions(self) -> None:
        self.start_captions(redo=True)

    def caption_progress(self, done: int, total: int) -> None:
        if self.caption_state[0] in ("making", "done"):
            self.caption_state = ("done" if done >= total else "making", done, total, "")
        self.session.captions = self.caption_store.texts(self.pictures)
        self._refresh_windows()

    def caption_failed(self, error) -> None:
        _, done, total, _ = self.caption_state
        self.caption_state = ("failed", done, total, error.detail or error.kind)
        self._refresh_windows()

    def edit_caption(self, key: str, text: str) -> None:
        self.caption_store.edit(key, text)
        self.session.captions = self.caption_store.texts(self.pictures)

    def caption_view(self) -> tuple[str, str, str, bool, bool]:
        """캡션 진행 막대: (칩, 칩 색, 글, [멈추기] 보임, [다시 만들기] 보임)."""
        state, done, total, reason = self.caption_state
        if state == "off":
            return "캡션 꺼짐", "off", "켜면 그림 설명을 만들어 질의응답 근거로 씁니다.", False, False
        if state == "paused":
            return "캡션 멈춤", "bad", "API 키를 확인하세요. 키가 맞으면 다시 시작합니다.", False, False
        if state == "making":
            return "캡션 만드는 중", "warn", f"{done} / {total}", True, False
        if state == "failed":
            return "캡션 멈춤", "bad", f"캡션을 만들지 못했어요: {reason}", False, True
        if state == "stopped":
            return "캡션 멈춤", "warn", f"{done} / {total}", False, True
        return "캡션 완료", "ok", "", False, True

    def on_action(self, key: str) -> None:
        if key == "settings":
            self.open_settings()
        elif key == "ask":
            self.open_question()
        elif key == "read":
            self.session.toggle_read()
        elif key == "speed":
            self.orb.set_speed(self.session.cycle_speed())
        elif key == "quit":
            self.quit()

    def quit(self) -> None:
        """[끝내기]: 읽기 정지 → 미리 만들기 취소 → 연동 닫기 → 설정 저장 → 종료."""
        self._timer.stop()
        self._monitor_timer.stop()
        if self.pipeline is not None:
            self.pipeline.stop()
        if self.presynth is not None:
            self.presynth.close()
        if self.caption_worker is not None:
            self.caption_worker.close()
        for w in (self.settings_window, self.script_window):
            if w is not None:
                w.close()
        self.session.shutdown()
        for close in self.closers:
            close()
        self.dog.close()
        self.orb.close()
        self.question.close()
        self.bubble.close()
        QApplication.quit()


def _make_app() -> QApplication:
    # 음높이를 지키며 배속하려면 FFmpeg 백엔드여야 한다(Windows 자체 백엔드는 못 한다)
    os.environ.setdefault("QT_MEDIA_BACKEND", "ffmpeg")
    app = QApplication.instance() or QApplication(sys.argv[:1])
    app.setQuitOnLastWindowClosed(False)  # 보통 창(설정 창)을 닫아도 강아지와 원 버튼은 남는다
    app.setStyleSheet(theme.STYLESHEET)
    platform_mac.become_accessory()
    return app


def run(path: str | None = None) -> int:
    from ..link import PowerPointLink
    from ..voice.cache import Presynth, VoiceCache
    from ..voice.engines import EngineUnavailable, make_engine
    from ..voice.pipeline import Pipeline, VoiceThread
    from ..voice.player import Player

    app = _make_app()
    invoker = MainThreadInvoker()
    settings = Settings()
    cache = VoiceCache(settings.folder)
    closers: list[Callable[[], None]] = []
    pipeline = None
    try:
        os_engine = make_engine(dict(settings["x5_factor"]))
        pipeline = Pipeline(os_engine, Player(), VoiceThread(invoker.post), lambda: cloud_engine(settings.data), cache)
        closers.append(pipeline.close)
    except EngineUnavailable as e:
        QMessageBox.critical(None, APP_TITLE, str(e))

    def link_factory(on_change, on_error):
        return PowerPointLink(invoker.wrap(on_change), invoker.wrap(on_error))

    def llm_start(provider, key, req, on_chunk, on_done, on_error):
        return start_stream(lambda: make_provider(provider, key), req, invoker.wrap(on_chunk),
                            invoker.wrap(on_done), invoker.wrap(on_error))

    session = Session(settings, pipeline, QtClock(), link_factory, llm_start=llm_start)
    if pipeline is not None:
        pipeline.on_done = session.voice_done
        pipeline.on_error = session.voice_error
        pipeline.on_cloud_failed = session.cloud_voice_failed
        pipeline.on_voice_chosen = session.cloud_voice_chosen
    controller = Controller(settings, session, pipeline, None, closers)
    controller.invoker = invoker
    controller.presynth = Presynth(cache, invoker.wrap(controller.presynth_progress),
                                   invoker.wrap(controller.presynth_failed))

    def caption_provider():
        return make_provider(settings["provider"], settings["keys"][settings["provider"]])

    controller.caption_worker = CaptionWorker(controller.caption_store, caption_provider,
                                              invoker.wrap(controller.caption_progress),
                                              invoker.wrap(controller.caption_failed))
    controller.start()
    if path:
        session.open_file(path)
    return app.exec()


# ---------------------------------------------------------------- --self-check

def self_check(out) -> int:
    """빌드 서버와 받는 사람 확인용. 줄마다 OK/FAIL을 쓰고, 모두 OK면 0을 돌려준다."""
    results: list[bool] = []

    def report(ok: bool, name: str, detail: str) -> None:
        results.append(ok)
        out.write(f"{'OK  ' if ok else 'FAIL'} {name}: {detail}\n")
        out.flush()

    if os.environ.get("QT_QPA_PLATFORM", "").startswith("offscreen"):
        del os.environ["QT_QPA_PLATFORM"]  # 실제 플랫폼으로 창을 만들어 봐야 한다
    app = _make_app()

    # 1. 강아지·원 버튼 창을 화면 밖 좌표에 한 번 만든다
    orb = None
    try:
        dog, orb = DogWindow(), OrbWindow(lambda _key: None)
        for w in (dog, orb):
            w.move(-4000, -4000)
            w.show()
        deadline = time.monotonic() + 1.0
        while time.monotonic() < deadline:
            app.processEvents()
            time.sleep(0.02)
        shown = dog.isVisible() and orb.isVisible() and bool(dog.winId()) and bool(orb.winId())
        platform = QGuiApplication.platformName()
        report(shown and platform != "offscreen", "화면", f"Qt {qVersion()}, 플랫폼 {platform}, 강아지·원 버튼 창")
    except Exception as e:
        report(False, "화면", f"{type(e).__name__}: {e}")

    # 2. 스프라이트 세 장
    try:
        bad = [d for d in DOGS if len(load_sheet(d)) != SWEAT_ROW + 1
               or any(f.isNull() for row in load_sheet(d) for f in row[:FRAMES])]
        report(not bad, "스프라이트", "세 장 읽음" if not bad else f"못 읽음: {', '.join(bad)}")
    except Exception as e:
        report(False, "스프라이트", f"{type(e).__name__}: {e}")

    # 3. 기본 음성으로 한 문장 합성(재생하지 않음). 실제처럼 음성 스레드에서 한다
    report(*_check_synth())

    # 4. 설정 폴더
    try:
        folder = app_folder()
        folder.mkdir(parents=True, exist_ok=True)
        report(folder.is_dir(), "설정 폴더", str(folder))
    except OSError as e:
        report(False, "설정 폴더", f"{e}")

    # 5. macOS 보조
    if sys.platform == "darwin":
        ok, detail = platform_mac.check(orb) if orb is not None else (False, "창 없음")
        report(ok, "macOS 보조", detail)
    else:
        report(True, "macOS 보조", "해당 없음")

    # 6. 재생 준비: 소리는 내지 않고, FFmpeg 백엔드가 올라와 음높이를 지키며 늘일 수 있는지만 본다
    report(*_check_playback())

    return 0 if all(results) else 1


def _check_playback() -> tuple[bool, str, str]:
    try:
        from PySide6.QtMultimedia import QMediaDevices, QMediaPlayer

        player = QMediaPlayer()
        availability = player.pitchCompensationAvailability()
        ok = availability != QMediaPlayer.PitchCompensationAvailability.Unavailable
        outputs = len(QMediaDevices.audioOutputs())
        return ok, "재생 준비", f"음높이 유지 {availability.name}, 출력 장치 {outputs}개"
    except Exception as e:
        return False, "재생 준비", f"{type(e).__name__}: {e}"


def _check_synth() -> tuple[bool, str, str]:
    from ..voice.engines import make_engine

    result: dict = {}

    def work():
        engine = None
        try:
            engine = make_engine()
            voices = engine.voices()
            voice = next((v for v in voices if v.is_korean), None)
            text = "안녕하세요. 자가 점검입니다." if voice else "This is a self check."
            synth = engine.synth(text, voice.id if voice else None, 1.0)
            result["ok"] = len(synth.pcm) > 0
            result["detail"] = (f"{voice.name if voice else '기본 음성'}, {synth.seconds:.2f}초, "
                                f"낱말 {len(synth.words)}개")
        except Exception as e:
            result["ok"], result["detail"] = False, f"{type(e).__name__}: {e}"
        finally:
            if engine is not None:
                engine.close()

    worker = threading.Thread(target=work, name="voice-check", daemon=True)
    worker.start()
    worker.join(60)
    if worker.is_alive():
        return False, "음성 합성", "60초 안에 끝나지 않음"
    return bool(result.get("ok")), "음성 합성", result.get("detail", "")

